"""Validador de pagarés del fondo de ahorro (aplicación de escritorio).

Compara un reporte de Excel contra pagarés en PDF y marca las celdas:
verde = coincide, rojo = no coincide, amarillo = no está en el PDF.
En Aval #1 y Aval #2 verifica si hay tinta (firma) sobre la línea FIRMA.
"""
import os
import queue
import re
import sys
import threading
import unicodedata
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

try:
    import pymupdf as fitz  # PyMuPDF (nombre actual)
except ImportError:
    import fitz
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

FILL = {k: PatternFill("solid", fgColor=c)
        for k, c in {"ok": "C6EFCE", "bad": "FFC7CE", "warn": "FFEB9C"}.items()}
COLS = {"id": "IDCOLABORADOR", "nombre": "NOMBRE", "quincena": "QUINCENA",
        "capital": "CAPITAL", "monto": "MONTOQUINCENAL", "interes": "INTERES",
        "aval1": "AVAL#1", "aval2": "AVAL#2"}


def money(s):
    try:
        return round(float(str(s).replace(",", "").replace("$", "")), 2)
    except (TypeError, ValueError):
        return None


def norm(s):
    s = unicodedata.normalize("NFD", "" if s is None else str(s))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", "", s).upper()


def tinta(page, rects, dpi=144, umbral=140, minimo=25):
    """True si hay más de `minimo` píxeles oscuros en las zonas indicadas."""
    n = 0
    for r in rects:
        pix = page.get_pixmap(clip=r, dpi=dpi, colorspace=fitz.csGRAY)
        n += sum(1 for b in pix.samples if b < umbral)
    return n > minimo


def leer_pdf(ruta):
    doc = fitz.open(ruta)
    texto = "\n".join(p.get_text() for p in doc)
    nombres = re.findall(r"NOMBRE\s*:\s*(.+?)\s+NUM\. COLAB\.\s*:\s*(\d+)", texto)

    def buscar(rx):
        m = re.search(rx, texto)
        return money(m.group(1)) if m else None

    q = re.search(r"([\d,]+\.\d\d)[^\w*]+([\d*]+)[^\w*]+\d\d/\d\d/\d{4}", texto)
    datos = {
        "id": int(nombres[0][1]) if nombres else None,
        "nombre": nombres[0][0] if nombres else None,
        "capital": buscar(r"IMPORTE PRESTAMO\s*\$?\s*([\d,]+\.\d+)"),
        "monto": buscar(r"DESCUENTO QUINCENAL\s*\$?\s*([\d,]+\.\d+)"),
        "interes": buscar(r"INTERESES\s*\$?\s*([\d,]+\.\d+)"),
        "quincena": q.group(2) if q else None,
        "firmas": [],
    }
    y_min = None
    for pagina in doc:
        hit = pagina.search_for("DATOS DE LOS AVALES")
        if hit:
            y_min = hit[0].y1
        if y_min is None:
            continue
        palabras = pagina.get_text("words")  # x0, y0, x1, y1, texto, ...
        firmas = sorted((w for w in palabras if w[4].startswith("FIRMA") and w[1] > y_min),
                        key=lambda w: w[1])
        for w in firmas:
            subr = sorted((u for u in palabras if u[4].startswith("__")
                           and abs(u[3] - w[3]) < 3 and u[0] > w[0]), key=lambda u: u[0])
            x0, x1, base = (subr[0][0], subr[0][2], subr[0][3]) if subr else (w[2] + 30, w[2] + 130, w[3])
            zonas = [fitz.Rect(x0 + 6, base - 18, x1, base - 4),   # sobre la línea
                     fitz.Rect(x0 + 6, base + 1, x1, base + 12)]   # justo debajo
            datos["firmas"].append(tinta(pagina, zonas))
        y_min = -1  # en páginas siguientes se revisa desde arriba
    return datos


def validar_excel(xlsx, datos, salida):
    wb = load_workbook(xlsx)
    ws = wb.worksheets[0]
    enc = next((r for r in range(1, 11)
                if any(norm(c.value) == COLS["id"] for c in ws[r])), None)
    if enc is None:
        raise ValueError("No encontré la columna 'ID Colaborador' en el Excel.")
    col = {norm(c.value): c.column for c in ws[enc] if c.value is not None}
    faltan = [v for v in COLS.values() if v not in col]
    if faltan:
        raise ValueError("Faltan columnas en el Excel: " + ", ".join(faltan))
    cont, usados = {"ok": 0, "bad": 0, "warn": 0}, set()
    for r in range(enc + 1, ws.max_row + 1):
        v = ws.cell(r, col[COLS["id"]]).value
        try:
            idv = int(str(v).strip())
        except (TypeError, ValueError):
            continue
        d = datos.get(idv)
        usados.add(idv) if d else None

        def celda(k):
            return ws.cell(r, col[COLS[k]])

        def marca(k, ok):
            est = "warn" if (d is None or ok is None) else ("ok" if ok else "bad")
            celda(k).fill = FILL[est]
            cont[est] += 1

        def cmp(k, f):
            return None if (d is None or d[k] is None) else f(celda(k).value) == f(d[k])

        marca("id", True if d else None)
        marca("nombre", cmp("nombre", norm))
        marca("capital", cmp("capital", money))
        marca("monto", cmp("monto", money))
        marca("interes", cmp("interes", money))
        marca("quincena", cmp("quincena", lambda x: str(x).strip()))
        for i in (0, 1):
            f = d["firmas"][i] if d and len(d["firmas"]) > i else None
            marca(f"aval{i + 1}", f)
    wb.save(salida)
    return cont, [k for k in datos if k not in usados]


def recurso(nombre):
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, nombre)


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Validador de pagarés | The Palace Company")
        self.geometry("620x380")
        self.xlsx = self.carpeta = None
        self.cola = queue.Queue()
        top = tk.Frame(self, bg="#14213d")
        top.pack(fill="x")
        try:
            img = tk.PhotoImage(file=recurso("logo.png"))
            img = img.subsample(max(1, -(-img.height() // 64)))
            tk.Label(top, image=img, bg="#14213d").pack(side="left", padx=12, pady=8)
            self.img = img
        except tk.TclError:
            tk.Label(top, text="TPC", fg="#e3c778", bg="#14213d",
                     font=("Georgia", 20, "bold")).pack(side="left", padx=16, pady=8)
        tk.Label(top, text="Validador de pagarés\nFondo de ahorro de empleados", fg="white",
                 bg="#14213d", justify="left", font=("Georgia", 13)).pack(side="left")
        body = tk.Frame(self, padx=20, pady=16)
        body.pack(fill="both", expand=True)
        self.l1 = tk.Label(body, text="Excel: sin seleccionar", anchor="w")
        self.l2 = tk.Label(body, text="Carpeta de PDF: sin seleccionar", anchor="w")
        ttk.Button(body, text="1. Seleccionar Excel", command=self.pedir_xlsx).pack(anchor="w")
        self.l1.pack(fill="x", pady=(2, 10))
        ttk.Button(body, text="2. Seleccionar carpeta de PDF", command=self.pedir_carpeta).pack(anchor="w")
        self.l2.pack(fill="x", pady=(2, 10))
        self.go = ttk.Button(body, text="3. Validar", command=self.iniciar, state="disabled")
        self.go.pack(anchor="w")
        self.barra = ttk.Progressbar(body, length=400)
        self.barra.pack(fill="x", pady=10)
        self.estado = tk.Label(body, text="", anchor="w", justify="left", wraplength=560)
        self.estado.pack(fill="x")
        self.after(150, self.revisar_cola)

    def pedir_xlsx(self):
        p = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx")])
        if p:
            self.xlsx = p
            self.l1.config(text="Excel: " + Path(p).name)
            self.habilitar()

    def pedir_carpeta(self):
        p = filedialog.askdirectory()
        if p:
            self.carpeta = p
            self.l2.config(text="Carpeta de PDF: " + p)
            self.habilitar()

    def habilitar(self):
        self.go.config(state="normal" if self.xlsx and self.carpeta else "disabled")

    def iniciar(self):
        self.go.config(state="disabled")
        threading.Thread(target=self.trabajo, daemon=True).start()

    def trabajo(self):
        try:
            pdfs = [p for p in Path(self.carpeta).rglob("*") if p.suffix.lower() == ".pdf"]
            if not pdfs:
                raise ValueError("No hay archivos PDF en la carpeta.")
            datos, errores = {}, []
            for i, p in enumerate(pdfs, 1):
                self.cola.put(("avance", i, len(pdfs)))
                try:
                    d = leer_pdf(str(p))
                    if d["id"] is None:
                        errores.append(p.name)
                    else:
                        datos[d["id"]] = d
                except Exception:
                    errores.append(p.name)
            salida = str(Path(self.xlsx).with_name(Path(self.xlsx).stem + "_validado.xlsx"))
            cont, sin_fila = validar_excel(self.xlsx, datos, salida)
            msg = (f"Listo. Verdes: {cont['ok']}  Rojas: {cont['bad']}  Amarillas: {cont['warn']}\n"
                   f"PDF sin fila en el Excel: {len(sin_fila)}  PDF no legibles: {len(errores)}\n"
                   f"Guardado en: {salida}")
            self.cola.put(("fin", msg))
        except PermissionError:
            self.cola.put(("error", "No pude guardar el archivo. Cierra el Excel validado si lo tienes abierto."))
        except Exception as e:
            self.cola.put(("error", str(e)))

    def revisar_cola(self):
        try:
            while True:
                m = self.cola.get_nowait()
                if m[0] == "avance":
                    self.barra.config(maximum=m[2], value=m[1])
                    self.estado.config(text=f"Leyendo PDF {m[1]} de {m[2]}")
                else:
                    self.estado.config(text=m[1])
                    self.go.config(state="normal")
                    if m[0] == "error":
                        messagebox.showerror("Error", m[1])
        except queue.Empty:
            pass
        self.after(150, self.revisar_cola)


if __name__ == "__main__":
    App().mainloop()
